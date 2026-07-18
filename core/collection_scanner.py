#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
collection_scanner.py — F03 : mode « scan de collection »
==========================================================

Session de détection continue qui déduplique les cartes vues et construit
un inventaire avec la valeur totale de la collection en fin de session.

Principe : à chaque frame, les détections (idéalement identifiées par F01,
`core/card_identifier.py`) sont accumulées par carte. Une carte est
**confirmée** — ajoutée une seule fois à l'inventaire — après avoir été vue
sur ``min_hits`` frames. Une même carte présentée 10 secondes n'apparaît
donc qu'une fois ; deux exemplaires posés côte à côte comptent pour une
quantité de 2 (maximum de détections simultanées).

Mode dégradé : sans identification F01 (index absent), la déduplication se
fait par classe YOLO (clé ``yolo:<classe>``) — moins précis, mais le scan
reste utilisable sur les cartes entraînées.

Les prix viennent de la base locale existante (models/cards_database.yaml,
alimentée par l'intégration TCGdex/Cardmarket) — aucun appel réseau.

Usage :
    >>> scanner = CollectionScanner()
    >>> for frame_detections in session:          # liste de Detection
    ...     scanner.observe_frame(frame_detections)
    >>> scanner.summary()['total_value']
    12.5
    >>> scanner.export_csv(); scanner.export_excel()
"""
import csv
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    from .utils import PATHS, safe_print
except ImportError:
    from utils import PATHS, safe_print

DEFAULT_OUTPUT_DIR = PATHS.get('directories', {}).get('output_collection_scans',
                                                      'output/collection_scans')

# Colonnes de l'export (CSV et Excel)
EXPORT_COLUMNS = [
    "card_id", "name", "set_id", "local_id", "quantity",
    "price", "price_max", "value", "value_max",
    "best_score", "hits", "first_seen", "last_seen",
    "condition",
]


@dataclass
class ScannedCard:
    """Une carte confirmée de l'inventaire."""
    key: str                        # clé de dédup (card_id ou yolo:<classe>)
    card_id: Optional[str]          # ex: "swsh7_003" (None en mode dégradé)
    name: str                       # nom affichable
    set_id: Optional[str] = None
    local_id: Optional[str] = None
    quantity: int = 1               # max de détections simultanées
    hits: int = 0                   # frames où la carte a été vue
    best_score: Optional[float] = None
    price: Optional[float] = None
    price_max: Optional[float] = None
    first_seen: float = 0.0         # time.time()
    last_seen: float = 0.0
    # Grading F02 (si actif pendant le scan) : meilleur état observé —
    # la valeur est pondérée par le facteur correspondant
    condition: Optional[str] = None          # NM / EX / GD / PL
    condition_score: Optional[float] = None
    price_factor: Optional[float] = None

    @property
    def value(self) -> Optional[float]:
        if self.price is None:
            return None
        return self.price * self.quantity * (self.price_factor or 1.0)

    @property
    def value_max(self) -> Optional[float]:
        if self.price_max is None:
            return None
        return self.price_max * self.quantity * (self.price_factor or 1.0)


@dataclass
class _Track:
    """Accumulateur d'une carte pas encore confirmée."""
    hits: int = 0
    max_simultaneous: int = 0
    best_score: Optional[float] = None
    name: Optional[str] = None
    card_id: Optional[str] = None
    first_seen: float = field(default_factory=time.time)
    last_seen: float = field(default_factory=time.time)
    condition: Optional[str] = None
    condition_score: Optional[float] = None
    price_factor: Optional[float] = None


class CollectionScanner:
    """
    Accumule les détections frame par frame, déduplique et agrège
    l'inventaire de la session.

    Args:
        min_hits: frames où une carte doit être vue avant confirmation
        prices: {card_id: {'name', 'price', 'price_max'}} — None charge la
                base locale (models/cards_database.yaml), {} désactive les prix
        track_ttl_s: durée sans détection au-delà de laquelle un track non
                     confirmé est abandonné (évite que des faux positifs
                     espacés sur toute la session finissent « confirmés »)
    """

    def __init__(self, min_hits: int = 3, prices: Optional[Dict] = None,
                 track_ttl_s: float = 10.0):
        if min_hits < 1:
            raise ValueError("min_hits doit être >= 1")
        self.min_hits = min_hits
        self.track_ttl_s = track_ttl_s
        if prices is None:
            try:
                try:
                    from .price_cache import load_prices_with_cache
                except ImportError:
                    from price_cache import load_prices_with_cache
                prices = load_prices_with_cache()
            except Exception as e:
                safe_print(f"⚠️ Prix indisponibles: {e}")
                prices = {}
        self._prices = prices or {}
        self.started_at = time.time()
        self._tracks: Dict[str, _Track] = {}
        self._inventory: Dict[str, ScannedCard] = {}
        self.frames_seen = 0

    # ---------- Observation ----------

    @staticmethod
    def _observation_key(det: Any) -> Optional[str]:
        """Clé de dédup d'une détection : card_id F01, sinon classe YOLO."""
        card_id = getattr(det, 'card_id', None)
        if card_id:
            return card_id
        class_name = getattr(det, 'class_name', None)
        if class_name:
            return f"yolo:{class_name}"
        return None

    def observe_frame(self, detections: List[Any],
                      timestamp: Optional[float] = None) -> List[ScannedCard]:
        """
        Enregistre les détections d'UNE frame.

        Args:
            detections: objets avec les attributs de ``Detection``
                        (class_name, et si F01 actif : card_id, exact_name,
                        identify_score)
            timestamp: time.time() de la frame (défaut: maintenant)

        Returns:
            Liste des cartes NOUVELLEMENT confirmées par cette frame.
        """
        now = timestamp if timestamp is not None else time.time()
        self.frames_seen += 1

        # Expirer les tracks non confirmés trop anciens (faux positifs)
        stale = [key for key, track in self._tracks.items()
                 if now - track.last_seen > self.track_ttl_s]
        for key in stale:
            del self._tracks[key]

        # Regrouper la frame par clé (occurrences simultanées)
        frame_counts: Dict[str, List[Any]] = {}
        for det in detections:
            key = self._observation_key(det)
            if key:
                frame_counts.setdefault(key, []).append(det)

        confirmed_now = []
        for key, dets in frame_counts.items():
            best = max((getattr(d, 'identify_score', None) or 0.0)
                       for d in dets)
            det0 = dets[0]
            grade = self._best_grade(dets)

            if key in self._inventory:
                card = self._inventory[key]
                card.hits += 1
                card.last_seen = now
                card.quantity = max(card.quantity, len(dets))
                if best and (card.best_score is None or best > card.best_score):
                    card.best_score = best
                self._apply_grade(card, grade)
                continue

            track = self._tracks.setdefault(key, _Track(first_seen=now))
            track.hits += 1
            track.last_seen = now
            track.max_simultaneous = max(track.max_simultaneous, len(dets))
            if best and (track.best_score is None or best > track.best_score):
                track.best_score = best
            track.card_id = getattr(det0, 'card_id', None)
            track.name = (getattr(det0, 'exact_name', None)
                          or getattr(det0, 'class_name', None) or key)
            self._apply_grade(track, grade)

            if track.hits >= self.min_hits:
                card = self._confirm(key, track, now)
                confirmed_now.append(card)

        return confirmed_now

    @staticmethod
    def _best_grade(dets) -> Optional[tuple]:
        """Meilleur grading F02 d'une frame : (condition, score, facteur)."""
        best = None
        for d in dets:
            score = getattr(d, 'condition_score', None)
            if score is not None and (best is None or score > best[1]):
                best = (getattr(d, 'condition', None), score,
                        getattr(d, 'price_factor', None))
        return best

    @staticmethod
    def _apply_grade(target, grade: Optional[tuple]) -> None:
        """
        Retient le MEILLEUR état observé sur la session (les frames floues
        ou en biais sous-estiment l'état ; la meilleure vue est la plus
        proche de la réalité).
        """
        if grade is None:
            return
        condition, score, factor = grade
        if target.condition_score is None or score > target.condition_score:
            target.condition = condition
            target.condition_score = score
            target.price_factor = factor

    def _confirm(self, key: str, track: _Track, now: float) -> ScannedCard:
        """Promeut un track en carte d'inventaire (avec prix)."""
        card_id = track.card_id
        set_id = local_id = None
        if card_id and '_' in card_id:
            set_id, local_id = card_id.rsplit('_', 1)

        # Prix : par card_id F01, sinon via le mapping classe -> card_id
        price_key = card_id
        if price_key is None and key.startswith("yolo:"):
            try:
                from .card_mapping import get_card_id_from_class_name
            except ImportError:
                from card_mapping import get_card_id_from_class_name
            price_key = get_card_id_from_class_name(key[len("yolo:"):])
        info = self._prices.get(price_key, {}) if price_key else {}

        card = ScannedCard(
            key=key,
            card_id=card_id,
            name=info.get('name') or track.name or key,
            set_id=set_id,
            local_id=local_id,
            quantity=max(1, track.max_simultaneous),
            hits=track.hits,
            best_score=track.best_score,
            price=info.get('price'),
            price_max=info.get('price_max'),
            first_seen=track.first_seen,
            last_seen=now,
            condition=track.condition,
            condition_score=track.condition_score,
            price_factor=track.price_factor,
        )
        self._inventory[key] = card
        del self._tracks[key]
        return card

    # ---------- Inventaire & agrégats ----------

    @property
    def inventory(self) -> List[ScannedCard]:
        """Cartes confirmées, triées par première apparition."""
        return sorted(self._inventory.values(), key=lambda c: c.first_seen)

    def summary(self) -> Dict[str, Any]:
        """Agrégats de la session (pour le compteur live et le récap)."""
        cards = self.inventory
        priced = [c for c in cards if c.price is not None]
        return {
            "unique_cards": len(cards),
            "total_quantity": sum(c.quantity for c in cards),
            "total_value": round(sum(c.value for c in priced), 2),
            "total_value_max": round(sum(c.value_max or c.value for c in priced), 2),
            "unpriced_cards": len(cards) - len(priced),
            "frames_seen": self.frames_seen,
            "duration_s": round(time.time() - self.started_at, 1),
        }

    def reset(self) -> None:
        """Réinitialise la session (inventaire, tracks, compteurs)."""
        self._tracks.clear()
        self._inventory.clear()
        self.frames_seen = 0
        self.started_at = time.time()

    # ---------- Exports ----------

    def _default_path(self, ext: str) -> Path:
        stamp = time.strftime("%Y%m%d_%H%M%S", time.localtime(self.started_at))
        return Path(DEFAULT_OUTPUT_DIR) / f"scan_{stamp}.{ext}"

    def _rows(self) -> List[List[Any]]:
        rows = []
        for c in self.inventory:
            rows.append([
                c.card_id or c.key, c.name, c.set_id, c.local_id, c.quantity,
                c.price, c.price_max, c.value, c.value_max,
                round(c.best_score, 4) if c.best_score is not None else None,
                c.hits,
                time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(c.first_seen)),
                time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(c.last_seen)),
                c.condition,
            ])
        return rows

    def export_csv(self, path: Optional[str] = None) -> Path:
        """Exporte l'inventaire en CSV. Retourne le chemin écrit."""
        path = Path(path) if path else self._default_path("csv")
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(EXPORT_COLUMNS)
            writer.writerows(self._rows())
        return path

    def export_excel(self, path: Optional[str] = None) -> Optional[Path]:
        """
        Exporte l'inventaire en Excel (feuille Inventory + ligne de totaux).
        Retourne le chemin écrit, ou None si openpyxl n'est pas installé.
        """
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            safe_print("⚠️ openpyxl non installé — export Excel ignoré "
                       "(pip install openpyxl)")
            return None

        path = Path(path) if path else self._default_path("xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Inventory"
        bold = Font(bold=True)

        ws.append(EXPORT_COLUMNS)
        for cell in ws[1]:
            cell.font = bold
        for row in self._rows():
            ws.append(row)

        s = self.summary()
        totals = [""] * len(EXPORT_COLUMNS)
        totals[0] = "TOTAL"
        totals[1] = f"{s['unique_cards']} cartes uniques"
        totals[4] = s['total_quantity']
        totals[7] = s['total_value']
        totals[8] = s['total_value_max']
        ws.append(totals)
        for cell in ws[ws.max_row]:
            cell.font = bold

        # Largeurs de colonnes lisibles
        for col, width in zip("ABCDEFGHIJKLMN",
                              [14, 28, 10, 9, 9, 9, 10, 9, 10, 11, 6, 20, 20, 10]):
            ws.column_dimensions[col].width = width

        wb.save(path)
        return path
